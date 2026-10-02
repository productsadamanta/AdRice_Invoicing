"""Evaluate the small formula subset used by the report generators and write the results into the
workbook as cached values, so viewers that do not calculate (previews, phone, e-mail) show numbers.

Why this exists: openpyxl writes formulas without results, and the machine this project runs on has
no LibreOffice/Excel to recalculate them. Excel still recalculates everything on open
(fullCalcOnLoad is set by the generators); this module only fills in the values beforehand and
doubles as an independent check that every formula evaluates without errors.

Supported: SUM, SUMIFS, COUNTIFS, COUNT, IF, + - * /, cell/range references (also across sheets).
Criteria support equality plus the operators >, <, >=, <=, <>.
"""
import ast
import re
import shutil
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL

TOK = re.compile(
    r'("[^"]*")'
    r"|(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)"
    r"|(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!\$?([A-Z]+)\$?(\d+)"
    r"|\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)"
    r'|(?<![A-Za-z_"])\$?([A-Z]{1,2})\$?(\d+)(?![\d(])'
)


def _num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else 0


def _match(cell, crit):
    if isinstance(crit, str):
        m = re.match(r"^(<>|>=|<=|>|<)(.*)$", crit)
        if m:
            op, rhs = m.groups()
            try:
                rhs_v = float(rhs)
            except ValueError:
                rhs_v = rhs
            if op == "<>":
                return not _match(cell, rhs_v if not isinstance(rhs_v, float) else rhs_v)
            if not isinstance(cell, (int, float)) or isinstance(cell, bool) or not isinstance(rhs_v, float):
                return False
            return {">": cell > rhs_v, "<": cell < rhs_v, ">=": cell >= rhs_v, "<=": cell <= rhs_v}[op]
        return isinstance(cell, str) and cell.lower() == crit.lower()
    return isinstance(cell, (int, float)) and not isinstance(cell, bool) and cell == crit


def f_sum(*args):
    t = 0
    for a in args:
        t += sum(_num(x) for x in a) if isinstance(a, list) else _num(a)
    return t


def f_sumifs(sumr, *pairs):
    t = 0
    for i, sv in enumerate(sumr):
        if all(_match(pairs[k][i], pairs[k + 1]) for k in range(0, len(pairs), 2)):
            t += _num(sv)
    return t


def f_countifs(*pairs):
    n = len(pairs[0])
    return sum(1 for i in range(n) if all(_match(pairs[k][i], pairs[k + 1]) for k in range(0, len(pairs), 2)))


def f_count(*args):
    return sum(1 for a in args for x in (a if isinstance(a, list) else [a])
               if isinstance(x, (int, float)) and not isinstance(x, bool))


class _LazyIf(ast.NodeTransformer):
    """Excel's IF only evaluates the branch it takes (e.g. IF(D=0,0,E/D) must not divide by zero)."""

    def visit_Call(self, node):
        self.generic_visit(node)
        if isinstance(node.func, ast.Name) and node.func.id == "IF" and len(node.args) == 3:
            lam = lambda e: ast.Lambda(args=ast.arguments(posonlyargs=[], args=[], kwonlyargs=[], kw_defaults=[], defaults=[]), body=e)  # noqa: E731
            node.args = [node.args[0], lam(node.args[1]), lam(node.args[2])]
        return node


class Book:
    def __init__(self, cells):
        self.cells = cells
        self.memo = {}

    def val(self, sheet, coord):
        key = (sheet, coord)
        if key in self.memo:
            return self.memo[key]
        v = self.cells.get(key)
        if isinstance(v, str) and v.startswith("="):
            v = self.ev(sheet, v[1:])
        self.memo[key] = v
        return v

    def rng(self, sheet, c1, r1, c2, r2):
        return [self.val(sheet, "%s%d" % (CL(ci), r))
                for r in range(int(r1), int(r2) + 1) for ci in range(CI(c1), CI(c2) + 1)]

    def ev(self, sheet, f):
        def rep(m):
            if m.group(1):
                return m.group(1)
            if m.group(4):
                return "R(%r,%r,%s,%r,%s)" % (m.group(2) or m.group(3), m.group(4), m.group(5), m.group(6), m.group(7))
            if m.group(10):
                return "V(%r,%r)" % (m.group(8) or m.group(9), m.group(10) + m.group(11))
            if m.group(12):
                return "R(%r,%r,%s,%r,%s)" % (sheet, m.group(12), m.group(13), m.group(14), m.group(15))
            return "V(%r,%r)" % (sheet, m.group(16) + m.group(17))

        py = TOK.sub(rep, f)
        py = re.sub(r"(?<![<>=!])=(?!=)", "==", py)
        env = dict(R=self.rng, V=self.val, SUM=f_sum, SUMIFS=f_sumifs, COUNTIFS=f_countifs, COUNT=f_count,
                   IF=lambda c, a, b: a() if c else b())
        tree = _LazyIf().visit(ast.parse(py, mode="eval"))
        ast.fix_missing_locations(tree)
        return eval(compile(tree, "<formula>", "eval"), dict(env, __builtins__={}))


def evaluate(path):
    """Return (workbook, Book, {(sheet, coord): value}, [errors])."""
    wb = load_workbook(path)
    cells = {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if c.value is not None}
    book = Book(cells)
    out, errs = {}, []
    for (sh, co), v in cells.items():
        if isinstance(v, str) and v.startswith("="):
            try:
                out[(sh, co)] = book.val(sh, co)
            except Exception as e:  # noqa: BLE001
                errs.append((sh, co, v, repr(e)))
    return wb, book, out, errs


def add_cached_values(path):
    """Write evaluated results into the formula cells of an openpyxl-generated .xlsx. Returns (written, total)."""
    _, _, out, errs = evaluate(path)
    if errs:
        raise RuntimeError("formula errors: %s" % errs[:3])
    z = zipfile.ZipFile(path)
    wbx = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid2file = {}
    for m in re.finditer(r"<Relationship ([^>]*)/>", rels):
        a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        rid2file[a["Id"]] = a["Target"]
    sheet2file = {}
    for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx):
        t = rid2file[m.group(2)].lstrip("/")
        sheet2file[m.group(1).replace("&amp;", "&")] = t if t.startswith("xl/") else "xl/" + t
    by_file = {}
    for (sh, co), v in out.items():
        by_file.setdefault(sheet2file[sh], {})[co] = v
    tmp = path + ".tmp"
    written = 0
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in by_file:
                vals = by_file[item.filename]

                def rep(m):
                    nonlocal written
                    v = vals.get(m.group(2))
                    if isinstance(v, bool) or not isinstance(v, (int, float)):
                        return m.group(0)
                    written += 1
                    return "%s%s<v>%s</v></c>" % (m.group(1), m.group(3), repr(float(v)) if isinstance(v, float) else v)

                data = re.sub(r'(<c r="([A-Z]+\d+)"[^>]*>)(<f>.*?</f>)<v\s*/></c>', rep, data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)
    return written, len(out)
