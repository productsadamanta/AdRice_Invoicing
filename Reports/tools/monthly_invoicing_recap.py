#!/usr/bin/env python
"""Monthly invoicing recap per advertiser: QuickBooks statement (PDF) -> invoices -> offers -> leads.

For every advertiser the workbook explains the statement total: which invoices are on it, and for each
invoice which AdRice offers (campaigns) and how many leads at what price, all taken from the Invoicing
Hub backup. The statement PDF is embedded in the advertiser's sheet as a picture.

Workbook: 'Summary' (all advertisers) + one sheet per advertiser (invoices, by country, offer recap,
invoice lines, statement picture). All totals are formulas; the Summary ties every advertiser back to
the total printed on its PDF.

Usage (from the Invoicing folder):
    python Reports/tools/monthly_invoicing_recap.py --month 2026-05 \
        --statements "trendi may.pdf" "tradeexpress may.pdf" "smartmedia may.pdf"
Optional: --backup <adrice_backup_*.json> (default: newest in ArchiveBackups), --out <file.xlsx>
"""
import argparse
import calendar
import collections
import datetime as dt
import glob
import io
import json
import math
import os
import re
import sys

import pdfplumber
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink
from PIL import Image, ImageChops

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import xlsx_formula_cache  # noqa: E402
from report_common import *  # noqa: E402,F401,F403

# ---------------------------------------------------------------- configuration
ADV_MATCH = {"TrendiSupply": r"TrendiSupply", "TradeExpress": r"TRADEXPERESS|TradeExpress",
             "SmartMediaSolving": r"SmartMediaSolving"}
ENTITY = {"TrendiSupply": "TrendiSupply Kft., Budapest, Hungary",
          "TradeExpress": "TRADEXPERESS PM LTD, Sofia, Bulgaria",
          "SmartMediaSolving": "SmartMediaSolving OÜ, Tallinn, Estonia"}
NOBREAK = "No breakdown (not in Hub)"


# ---------------------------------------------------------------- inputs
def load_db(path):
    d = json.load(open(path, encoding="utf-8"))
    return d.get("invoicingDB", d)


def newest_backup(folder):
    files = sorted(glob.glob(os.path.join(folder, "adrice_backup_*.json")))
    files = [f for f in files if not re.search(r"(reconstructed|REPAIRED|merged)", f, re.I)]
    return files[-1]


def parse_statement(path):
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        page = pdf.pages[0]
        im = page.to_image(resolution=130).original.convert("RGB")
    adv = next((a for a, pat in ADV_MATCH.items() if re.search(pat, text, re.I)), None)
    if not adv:
        raise SystemExit("Cannot tell which advertiser the statement belongs to: %s" % path)
    lines = []
    for m in re.finditer(r"(\d\d)/(\d\d)/(\d{4}) Invoice No\.(INV-?\d+) ([\d,]+\.\d\d) ([\d,]+\.\d\d)", text):
        d, mo, y, no, amt, rcv = m.groups()
        lines.append(dict(date=dt.date(int(y), int(mo), int(d)), no=re.sub(r"INV-?", "INV-", no),
                          amount=float(amt.replace(",", "")), received=float(rcv.replace(",", ""))))
    tot = re.search(r"EUR ([\d,]+\.\d\d)\s+EUR ([\d,]+\.\d\d)", text)
    no = re.search(r"STATEMENT NO\.\s*(\d+)", text)
    sd = re.search(r"DATE\s+(\d\d)/(\d\d)/(\d{4})", text)
    if not (lines and tot):
        raise SystemExit("Statement not parsed (no invoice lines or total): %s" % path)
    # crop whitespace and keep a small margin
    bbox = ImageChops.difference(im, Image.new("RGB", im.size, (255, 255, 255))).getbbox()
    x0, y0, x1, y1 = bbox
    im = im.crop((max(0, x0 - 12), max(0, y0 - 12), min(im.width, x1 + 12), min(im.height, y1 + 12)))
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return dict(adv=adv, path=path, lines=lines, total=float(tot.group(1).replace(",", "")),
                number=no.group(1) if no else "", date=dt.date(int(sd.group(3)), int(sd.group(2)), int(sd.group(1))) if sd else None,
                png=buf.getvalue(), size=im.size)


# ---------------------------------------------------------------- Hub data
class Hub:
    def __init__(self, db):
        self.db = db
        self.inv = {}      # invoice no -> (month key, log)
        self.dups = []
        self.title = {}    # offer id -> best title
        self.off_cc = collections.defaultdict(set)
        self.steps = collections.defaultdict(list)  # (month, offer, prev, curr) -> [(label, log id)]: counter steps already invoiced
        for mk, m in db.items():
            if not (isinstance(m, dict) and "logs" in m):
                continue
            for cc in COUNTRIES:
                for oid in (m.get(cc.lower()) or {}):
                    self.off_cc[oid].add(cc)
            for l in m["logs"]:
                no = l.get("invoiceNumber") or ""
                if no not in ("PRZENIESIENIE_DO_SALDA", "REKONCYLIACJA_DO_SALDA"):
                    for it in l.get("items", []):
                        if "prev" in it and it["delta"] != 0:
                            self.steps[(mk, str(it["id"]), it["prev"], it["curr"])].append((no, id(l)))
                if re.match(r"INV-?\d+$", no):
                    no = re.sub(r"INV-?", "INV-", no)
                    if no in self.inv:
                        self.dups.append(no)
                    else:
                        self.inv[no] = (mk, l)
                for it in l.get("items", []) + [i for a in l.get("adjustments", []) for i in a.get("items", [])]:
                    t = it.get("title", "")
                    if t.startswith("Lead Generation"):
                        self.title[str(it["id"])] = t

    def clean_title(self, oid, title):
        if not title.startswith("Lead Generation"):
            title = self.title.get(str(oid), title)
        return re.sub(r"^Lead Generation\s+", "", title)

    def country(self, oid, title, fallback=None):
        t = self.title.get(str(oid), title)
        hits = [cc for word, cc in TITLE_CC if re.search(r"\b" + word, t)]
        if hits:
            return hits[-1]
        c = self.off_cc.get(str(oid), set())
        if len(c) == 1:
            return next(iter(c))
        if fallback and fallback.upper() in COUNTRIES:
            return fallback.upper()
        return NA


def build_lines(st, hub, month, warnings):
    """Statement -> list of line dicts (one per offer on each invoice, plus corrections)."""
    adv, lines = st["adv"], []
    for sl in st["lines"]:
        h = hub.inv.get(sl["no"])
        if not h:
            warnings.append("%s: %s is on the statement but not in the Hub (amount %.2f) - added as a single line without breakdown." % (adv, sl["no"], sl["amount"]))
            lines.append(dict(inv=sl["no"], date=sl["date"], cat=NOBREAK, oid=NA, title="%s - no offer / lead breakdown (not in Hub)" % sl["no"],
                              cc=NA, leads=None, price=None, value=sl["amount"]))
            continue
        mk, log = h
        if TYPE2ADV.get(log["type"]) != adv:
            warnings.append("%s: %s is %s in the Hub (type %s)." % (adv, sl["no"], TYPE2ADV.get(log["type"]), log["type"]))
        tot = 0.0
        repeats = collections.defaultdict(lambda: [0, 0.0])  # earlier invoice -> [lines, value] whose counter step this invoice repeats
        for it in log["items"]:
            if it["delta"] == 0:
                continue
            tot += it["delta"] * it["price"]
            if "prev" in it:
                for other, log_id in hub.steps.get((mk, str(it["id"]), it["prev"], it["curr"]), []):
                    if log_id != id(log):  # same offer, same counter step already invoiced elsewhere
                        repeats[other][0] += 1
                        repeats[other][1] += it["delta"] * it["price"]
            lines.append(dict(inv=sl["no"], date=sl["date"], cat="%s leads" % month_label(mk), oid=offer_id(it["id"]),
                              title=hub.clean_title(it["id"], it["title"]), cc=hub.country(it["id"], it["title"]),
                              leads=it["delta"], price=it["price"], value=None, mk=mk))
        for a in log.get("adjustments", []):
            for it in a["items"]:
                tot += it["delta"] * it["price"]
                lines.append(dict(inv=sl["no"], date=sl["date"], cat="Correction for %s" % month_label(a["month"]), oid=offer_id(it["id"]),
                                  title=hub.clean_title(it["id"], it["title"]), cc=hub.country(it["id"], it["title"], a.get("account")),
                                  leads=it["delta"], price=it["price"], value=None, mk=a["month"]))
        if abs(tot - sl["amount"]) > 0.005:
            warnings.append("%s: %s Hub lines total %.2f but the statement says %.2f." % (adv, sl["no"], tot, sl["amount"]))
        for other, (n, val) in repeats.items():
            warnings.append("%s: %s has %d line(s) worth %s EUR with exactly the same counter step (same offer, same previous and new count) as %s - the same change may have been invoiced twice. Check both invoices." % (
                adv, sl["no"], n, "{:,.2f}".format(val), other))
    return lines


# ---------------------------------------------------------------- advertiser sheet
def advertiser_sheet(wb, adv, st, lines, month, warnings, notes=None):
    ws = wb.create_sheet(adv)
    # left block A:H (invoices, countries, offers), spacer I, right block J:R (picture + invoice lines)
    widths = dict(A=12, B=50, C=16, D=11, E=16, F=16, G=14, H=34, I=3, J=13, K=12, L=30, M=9, N=62, O=9, P=9, Q=10, R=14)
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    ml = month_label(month)
    put(ws, 1, 1, "%s - invoices dated %s" % (adv, ml), font=f_title)
    put(ws, 2, 1, "%s | Statement no. %s dated %s | Amounts in EUR" % (ENTITY.get(adv, adv), st["number"], st["date"].strftime("%d.%m.%Y") if st["date"] else ""), font=f_sub)

    if notes:  # reviewed explanations (annotations file), shown right under the title
        paragraphs = ["NOTE: " + notes[0]] + list(notes[1:])
        put(ws, 4, 1, "\n".join(paragraphs), wrap=True, fill=PatternFill("solid", fgColor="FFF2CC"))
        ws.merge_cells("A4:H4")
        ws.row_dimensions[4].height = 14 * sum(math.ceil(len(p) / 150) for p in paragraphs) + 8

    # ---- right block geometry: picture, then invoice lines
    img_w = 600
    img_h = int(st["size"][1] * img_w / st["size"][0])
    pic_rows = math.ceil(img_h / 20) + 1
    sec_row = 6 + pic_rows + 1             # 'Invoice lines' section header
    hdr_row = sec_row + 1
    first = hdr_row + 1
    last = first + len(lines) - 1

    def R(col):
        """Range of one invoice-line column. Callers pass the letter of the line columns in the original layout
        (I=invoice, K=category, L=offer id, N=country, O=leads, Q=value); the right block now starts one column further right."""
        c = chr(ord(col) + 1)
        return "$%s$%d:$%s$%d" % (c, first, c, last)

    # ---- invoice lines (right block, columns J:R)
    put(ws, 5, 10, "Statement from QuickBooks (PDF embedded as picture)", font=f_sec)
    img = XLImage(io.BytesIO(st["png"]))
    img.width, img.height = img_w, img_h
    ws.add_image(img, "J6")
    put(ws, sec_row, 10, "4. Invoice lines - leads recap detail (one row per offer on each invoice)", font=f_sec)
    header(ws, hdr_row, 10, ["Invoice no.", "Invoice date", "Category", "Offer ID", "Offer", "Country", "Leads", "Price (€)", "Value (€)"])
    for i, ln in enumerate(lines):
        r = first + i
        put(ws, r, 10, ln["inv"]); put(ws, r, 11, ln["date"], fmt=DATE); put(ws, r, 12, ln["cat"])
        put(ws, r, 13, ln["oid"]); put(ws, r, 14, ln["title"]); put(ws, r, 15, ln["cc"])
        if ln["leads"] is None:  # no breakdown: amount comes from the statement
            put(ws, r, 18, ln["value"], font=f_in, fmt=EUR, fill=fill_warn)
            ws.cell(r, 14).fill = fill_warn
        else:
            put(ws, r, 16, ln["leads"], font=f_in, fmt=INT); put(ws, r, 17, ln["price"], font=f_in, fmt=EUR)
            put(ws, r, 18, "=P%d*Q%d" % (r, r), fmt=EUR)
    ws.auto_filter.ref = "J%d:R%d" % (hdr_row, last)

    # ---- 1. invoices
    r = 5
    put(ws, r, 1, "1. Invoices on the statement (invoice date in %s)" % ml, font=f_sec)
    header(ws, r + 1, 1, ["Invoice no.", "Invoice date", "Lead period", "Leads", "Sum of invoice lines (€)", "Statement amount (€)", "Difference (€)",
                          "o/w credit lines = lead reductions (€)"])
    inv_first = r + 2
    for k, sl in enumerate(st["lines"]):
        rr = inv_first + k
        hub = [ln for ln in lines if ln["inv"] == sl["no"]]
        period = "n/a - not in Hub" if hub and hub[0]["cat"] == NOBREAK else month_label(hub[0]["mk"]) if hub and "mk" in hub[0] else ""
        put(ws, rr, 1, sl["no"]); put(ws, rr, 2, sl["date"], fmt=DATE, align="left"); put(ws, rr, 3, period)
        put(ws, rr, 4, "=SUMIFS(%s,%s,$A%d)" % (R("O"), R("I"), rr), fmt=INT)
        put(ws, rr, 5, "=SUMIFS(%s,%s,$A%d)" % (R("Q"), R("I"), rr), fmt=EUR)
        put(ws, rr, 6, sl["amount"], font=f_in, fmt=EUR)
        put(ws, rr, 7, "=E%d-F%d" % (rr, rr), fmt=EUR)
        put(ws, rr, 8, '=SUMIFS(%s,%s,$A%d,%s,"<0")' % (R("Q"), R("I"), rr, R("Q")), fmt=EUR)
        if period.startswith("n/a"):
            for c in (1, 3):
                ws.cell(rr, c).fill = fill_warn
    inv_last = inv_first + len(st["lines"]) - 1
    inv_tot = inv_last + 1
    put(ws, inv_tot, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in (2, 3):
        put(ws, inv_tot, c, None, fill=fill_tot)
    for c, fm in ((4, INT), (5, EUR), (6, EUR), (7, EUR), (8, EUR)):
        put(ws, inv_tot, c, "=SUM(%s%d:%s%d)" % (L(c), inv_first, L(c), inv_last), fmt=fm, bold=True, fill=fill_tot)
    pdf_row = inv_tot + 1
    put(ws, pdf_row, 1, "Total printed on the PDF")
    put(ws, pdf_row, 6, st["total"], font=f_in, fmt=EUR)
    put(ws, pdf_row, 7, "=F%d-F%d" % (inv_tot, pdf_row), fmt=EUR)
    put(ws, pdf_row, 2, "(check: invoices listed vs. total on the statement)", font=f_sub)

    # ---- 2. by country
    r = pdf_row + 2
    put(ws, r, 1, "2. By country", font=f_sec)
    header(ws, r + 1, 1, ["Code", "Country", "", "Leads", "Value (€)"])
    ccs = sorted({ln["cc"] for ln in lines}, key=lambda c: COUNTRIES.index(c) if c in COUNTRIES else 99)
    cc_first = r + 2
    for k, cc in enumerate(ccs):
        rr = cc_first + k
        put(ws, rr, 1, cc); put(ws, rr, 2, COUNTRY_NAME.get(cc, cc))
        put(ws, rr, 4, "=SUMIFS(%s,%s,$A%d)" % (R("O"), R("N"), rr), fmt=INT)
        put(ws, rr, 5, "=SUMIFS(%s,%s,$A%d)" % (R("Q"), R("N"), rr), fmt=EUR)
    cc_last = cc_first + len(ccs) - 1
    cc_tot = cc_last + 1
    put(ws, cc_tot, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in (2, 3):
        put(ws, cc_tot, c, None, fill=fill_tot)
    for c, fm in ((4, INT), (5, EUR)):
        put(ws, cc_tot, c, "=SUM(%s%d:%s%d)" % (L(c), cc_first, L(c), cc_last), fmt=fm, bold=True, fill=fill_tot)

    # ---- 3. offers (campaign recap)
    r = cc_tot + 2
    off_sec = r
    put(ws, r, 1, "3. Campaign recap - leads and value per AdRice offer", font=f_sec)
    header(ws, r + 1, 1, ["Offer ID", "Offer", "Country", "Leads", "Value (€)", "Avg price (€/lead)", "Share of total", "Note"])
    agg = collections.OrderedDict()
    for ln in lines:
        k = (ln["oid"], ln["cc"])
        a = agg.setdefault(k, dict(title=ln["title"], v=0.0, n=0, cr=[]))
        a["v"] += ln["value"] if ln["leads"] is None else ln["leads"] * ln["price"]
        a["n"] += ln["leads"] or 0
        if (ln["leads"] or 0) < 0 and ln["inv"] not in a["cr"]:
            a["cr"].append(ln["inv"])  # invoices carrying a lead reduction for this offer
    keys = sorted(agg, key=lambda k: (COUNTRIES.index(k[1]) if k[1] in COUNTRIES else 99, -agg[k]["v"], str(k[0])))
    of_first = r + 2
    of_last = of_first + len(keys) - 1
    of_tot = of_last + 1
    for k, key in enumerate(keys):
        rr = of_first + k
        put(ws, rr, 1, key[0]); put(ws, rr, 2, agg[key]["title"]); put(ws, rr, 3, key[1])
        put(ws, rr, 4, "=SUMIFS(%s,%s,$A%d,%s,$C%d)" % (R("O"), R("L"), rr, R("N"), rr), fmt=INT)
        put(ws, rr, 5, "=SUMIFS(%s,%s,$A%d,%s,$C%d)" % (R("Q"), R("L"), rr, R("N"), rr), fmt=EUR)
        put(ws, rr, 6, "=IF(D%d=0,0,E%d/D%d)" % (rr, rr, rr), fmt=EUR)
        put(ws, rr, 7, "=E%d/$E$%d" % (rr, of_tot), fmt=PCT)
        if key[0] == NA:
            put(ws, rr, 8, "No breakdown available (amount from the statement)", font=f_sub)
            for c in (1, 2, 3):
                ws.cell(rr, c).fill = fill_warn
        elif agg[key]["cr"]:
            net_credit = agg[key]["n"] < 0 or agg[key]["v"] < 0
            put(ws, rr, 8, ("Net credit: lead reduction on %s" if net_credit else "Includes a lead reduction on %s") % ", ".join(agg[key]["cr"]), font=f_sub)
    put(ws, of_tot, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in (2, 3, 6):
        put(ws, of_tot, c, None, fill=fill_tot)
    for c, fm in ((4, INT), (5, EUR)):
        put(ws, of_tot, c, "=SUM(%s%d:%s%d)" % (L(c), of_first, L(c), of_last), fmt=fm, bold=True, fill=fill_tot)
    put(ws, of_tot, 7, "=SUM(G%d:G%d)" % (of_first, of_last), fmt=PCT, bold=True, fill=fill_tot)
    put(ws, of_tot + 1, 1, "Difference: offer recap total - statement total")
    put(ws, of_tot + 1, 5, "=E%d-F%d" % (of_tot, inv_tot), fmt=EUR)

    # ---- navigation
    put(ws, 3, 1, "Jump to:", font=f_sub)
    link(ws, 3, 2, "1 Invoices", adv, "A5")
    link(ws, 3, 3, "2 By country", adv, "A%d" % (pdf_row + 2))
    link(ws, 3, 4, "3 Offers", adv, "A%d" % off_sec)
    link(ws, 3, 5, "4 Invoice lines", adv, "J%d" % sec_row)
    link(ws, 3, 6, "Statement", adv, "J5")
    ws.sheet_view.zoomScale = 90
    return dict(ws=ws, rng=lambda col: "'%s'!$%s$%d:$%s$%d" % (adv, chr(ord(col) + 1), first, chr(ord(col) + 1), last), inv_tot=inv_tot, pdf_row=pdf_row,
                of_tot=of_tot, of_first=of_first, of_last=of_last, ccs=ccs, cc_first=cc_first, cc_tot=cc_tot, n_inv=len(st["lines"]),
                inv_first=inv_first, inv_last=inv_last)


# ---------------------------------------------------------------- summary
def summary_sheet(wb, month, sheets, statements, lines_by_adv, backup_name, warnings):
    ws = wb.create_sheet("Summary", 0)
    ml = month_label(month)
    for k, w in dict(A=22, B=16, C=16, D=16, E=18, F=18, G=18, H=18, I=16, J=14, K=12).items():
        ws.column_dimensions[k].width = w
    put(ws, 1, 1, "Monthly invoicing recap - %s" % ml, font=f_title)
    put(ws, 2, 1, "Invoices dated in %s per the QuickBooks statements, explained by AdRice offer (campaign) and leads. Source of the breakdown: Invoicing Hub backup %s. Amounts in EUR." % (ml, backup_name), font=f_sub)

    put(ws, 4, 1, "1. Statement tie-out", font=f_sec)
    header(ws, 5, 1, ["Advertiser", "Statement no.", "Statement date", "# invoices", "Total printed on PDF (€)", "Σ invoices on statement (€)",
                      "Explained by offers (€)", "o/w without offer breakdown (€)", "Difference PDF - offers (€)", "Leads", "# offers"], height=44)
    for k, adv in enumerate(ADVS):
        if adv not in sheets:
            continue
        s, st = sheets[adv], statements[adv]
        r = 6 + k
        q = "'%s'!" % adv
        link(ws, r, 1, adv, adv, "A1")
        put(ws, r, 2, st["number"]); put(ws, r, 3, st["date"], fmt=DATE, align="left")
        put(ws, r, 4, '=COUNTIFS(%s!$F$%d:$F$%d,">0")' % ("'%s'" % adv, s["inv_first"], s["inv_last"]), fmt=INT)
        put(ws, r, 5, "=%sF%d" % (q, s["pdf_row"]), font=f_xs, fmt=EUR)
        put(ws, r, 6, "=%sF%d" % (q, s["inv_tot"]), font=f_xs, fmt=EUR)
        put(ws, r, 7, "=%sE%d" % (q, s["of_tot"]), font=f_xs, fmt=EUR, bold=True, fill=fill_key)
        put(ws, r, 8, '=SUMIFS(%s,%s,"%s")' % (s["rng"]("Q"), s["rng"]("K"), NOBREAK), fmt=EUR)
        put(ws, r, 9, "=E%d-G%d" % (r, r), fmt=EUR)
        put(ws, r, 10, "=%sD%d" % (q, s["of_tot"]), font=f_xs, fmt=INT)
        put(ws, r, 11, "=COUNT(%s!$A$%d:$A$%d)" % ("'%s'" % adv, s["of_first"], s["of_last"]), fmt=INT)
    t = 6 + len(ADVS)
    put(ws, t, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in (2, 3):
        put(ws, t, c, None, fill=fill_tot)
    for c, fm in ((4, INT), (5, EUR), (6, EUR), (7, EUR), (8, EUR), (9, EUR), (10, INT), (11, INT)):
        put(ws, t, c, "=SUM(%s6:%s%d)" % (L(c), L(c), t - 1), fmt=fm, bold=True, fill=fill_tot)

    # ---- composition
    cats = {}
    for adv, lines in lines_by_adv.items():
        for ln in lines:
            if ln["cat"] == "%s leads" % ml:
                key = (0, "")
            elif ln["cat"].endswith(" leads"):
                key = (1, ln["mk"])
            elif ln["cat"].startswith("Correction"):
                key = (2, ln["mk"])
            else:
                key = (3, "")
            cats[ln["cat"]] = key
    cat_list = sorted(cats, key=lambda c: cats[c])
    r0 = t + 2
    put(ws, r0, 1, "2. How each statement total is composed (by category of invoice line)", font=f_sec)
    header(ws, r0 + 1, 1, ["Advertiser"] + cat_list + ["Total (€)", "Check vs. PDF total (€)"], height=44)
    for k, adv in enumerate(ADVS):
        if adv not in sheets:
            continue
        s = sheets[adv]
        r = r0 + 2 + k
        put(ws, r, 1, adv, bold=True)
        for j, cat in enumerate(cat_list):
            put(ws, r, 2 + j, "=SUMIFS(%s,%s,%s$%d)" % (s["rng"]("Q"), s["rng"]("K"), L(2 + j), r0 + 1), fmt=EUR)
        tc = 2 + len(cat_list)
        put(ws, r, tc, "=SUM(B%d:%s%d)" % (r, L(tc - 1), r), fmt=EUR, bold=True, fill=fill_key)
        put(ws, r, tc + 1, "=%s%d-E%d" % (L(tc), r, 6 + k), fmt=EUR)
    rt = r0 + 2 + len(ADVS)
    put(ws, rt, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in range(2, 2 + len(cat_list) + 2):
        put(ws, rt, c, "=SUM(%s%d:%s%d)" % (L(c), r0 + 2, L(c), rt - 1), fmt=EUR, bold=True, fill=fill_tot)

    # ---- by company and country
    r1 = rt + 2
    put(ws, r1, 1, "3. Leads and value by advertiser and country", font=f_sec)
    header(ws, r1 + 1, 1, ["Advertiser", "Country", "Leads", "Value (€)"])
    r = r1 + 2
    sub_rows = []
    for adv in ADVS:
        if adv not in sheets:
            continue
        s = sheets[adv]
        first_r = r
        for cc in s["ccs"]:
            put(ws, r, 1, adv); put(ws, r, 2, "%s - %s" % (cc, COUNTRY_NAME.get(cc, cc)))
            put(ws, r, 3, "=SUMIFS(%s,%s,\"%s\")" % (s["rng"]("O"), s["rng"]("N"), cc), fmt=INT)
            put(ws, r, 4, "=SUMIFS(%s,%s,\"%s\")" % (s["rng"]("Q"), s["rng"]("N"), cc), fmt=EUR)
            r += 1
        put(ws, r, 1, adv + " - subtotal", bold=True, fill=fill_tot)
        put(ws, r, 2, None, fill=fill_tot)
        for c, fm in ((3, INT), (4, EUR)):
            put(ws, r, c, "=SUM(%s%d:%s%d)" % (L(c), first_r, L(c), r - 1), fmt=fm, bold=True, fill=fill_tot)
        sub_rows.append(r)
        r += 1
    put(ws, r, 1, "GRAND TOTAL", bold=True, fill=fill_key)
    put(ws, r, 2, None, fill=fill_key)
    for c, fm in ((3, INT), (4, EUR)):
        put(ws, r, c, "=" + "+".join("%s%d" % (L(c), x) for x in sub_rows), fmt=fm, bold=True, fill=fill_key)

    # ---- notes
    r += 2
    put(ws, r, 1, "Notes", font=f_sec)
    notes = [
        "Basis: invoices with an invoice date in %s, exactly as listed on the QuickBooks statements (all shown as fully received). The statements are dated %s." % (ml, ", ".join(sorted({st['date'].strftime('%d.%m.%Y') for st in statements.values() if st['date']}))),
        "Each advertiser sheet lists the invoices, then the leads and value per AdRice offer (campaign), then every invoice line behind those numbers, and carries the statement PDF as a picture.",
        "Lead period: an invoice dated in %s can bill leads of the previous month (e.g. the first invoices of the month) and may carry corrections of earlier leads; both are labelled in the 'Category' column and in table 2." % ml,
        "Leads = quantities on the invoice lines, including the quantities in corrections. Value = leads x price per line, as recorded in the Hub when the invoice was generated.",
        "Negative numbers (red, with a minus sign) are credits: leads that were invoiced earlier and later reduced. They lower the invoice total. An offer can show a negative net in the month when its original leads were invoiced in an earlier period (column 'o/w credit lines' on each advertiser sheet, note next to the offer).",
        "Blue = values taken from the statements or the Hub, black = formulas, green = links to another sheet, yellow = lines that have no Hub breakdown (amount taken from the statement).",
    ]
    notes.extend(warnings)  # already prefixed ('ATTENTION: ...' or 'Checked: ...') by main()
    for n in notes:
        r += 1
        put(ws, r, 1, "•")
        put(ws, r, 2, n, wrap=True, fill=fill_warn if n.startswith("ATTENTION") else None)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
        ws.row_dimensions[r].height = 30 if len(n) < 190 else 44
    ws.sheet_view.zoomScale = 95
    return ws


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--month", required=True, help="Report month, YYYY-MM (invoice date)")
    ap.add_argument("--statements", nargs="+", required=True, help="QuickBooks statement PDFs, one per advertiser")
    ap.add_argument("--backup", help="Hub backup JSON (default: newest in ArchiveBackups)")
    ap.add_argument("--out", help="Output .xlsx")
    ap.add_argument("--annotations", help="JSON with reviewed notes (default: Reports/annotations/<month>.json if it exists)")
    a = ap.parse_args()
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    ann_path = a.annotations or os.path.join(root, "Reports", "annotations", "%s.json" % a.month)
    ann = json.load(open(ann_path, encoding="utf-8")) if os.path.exists(ann_path) else {}
    backup = a.backup or newest_backup(os.path.join(root, "ArchiveBackups"))
    out = a.out or os.path.join(root, "Reports", "Monthly_Invoicing_Recap_%s.xlsx" % a.month)
    hub = Hub(load_db(backup))
    warnings = []
    statements, lines_by_adv = {}, {}
    for p in a.statements:
        st = parse_statement(p)
        off = [sl for sl in st["lines"] if sl["date"].strftime("%Y-%m") != a.month]
        if off:
            warnings.append("%s: %d statement line(s) are not dated in %s." % (st["adv"], len(off), month_label(a.month)))
        if abs(sum(sl["amount"] for sl in st["lines"]) - st["total"]) > 0.005:
            warnings.append("%s: lines on the PDF do not add up to its printed total." % st["adv"])
        statements[st["adv"]] = st
        lines_by_adv[st["adv"]] = build_lines(st, hub, a.month, warnings)
    on_statements = {sl["no"] for st in statements.values() for sl in st["lines"]}
    dups = sorted(set(hub.dups) & on_statements)
    if dups:
        warnings.append("Invoice numbers used more than once in the Hub (first one used): %s" % ", ".join(dups))

    # Warnings the reviewer has already checked are shown as 'Checked: <resolution>' instead of 'ATTENTION: ...'
    resolved = ann.get("resolved_warnings", {})
    summary_notes = []
    for w in warnings:
        hit = next((txt for key, txt in resolved.items() if key in w), None)
        summary_notes.append(("Checked: " + hit) if hit else ("ATTENTION: " + w))
    summary_notes.extend(ann.get("summary_notes", []))

    wb = Workbook()
    wb.remove(wb.active)
    sheets = {adv: advertiser_sheet(wb, adv, statements[adv], lines_by_adv[adv], a.month, warnings, ann.get("advertiser_notes", {}).get(adv))
              for adv in ADVS if adv in statements}
    summary_sheet(wb, a.month, sheets, statements, lines_by_adv, os.path.basename(backup), summary_notes)
    wb._sheets = [wb["Summary"]] + [wb[adv] for adv in ADVS if adv in sheets]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    written, total = xlsx_formula_cache.add_cached_values(out)
    print("saved", out)
    print("backup:", backup)
    print("cached values: %d of %d formulas" % (written, total))
    for adv in ADVS:
        if adv in statements:
            print("  %-18s lines %4d | PDF total %12s | Hub+statement lines %12s" % (
                adv, len(lines_by_adv[adv]), "{:,.2f}".format(statements[adv]["total"]),
                "{:,.2f}".format(sum((ln["value"] if ln["leads"] is None else ln["leads"] * ln["price"]) for ln in lines_by_adv[adv]))))
    for w in warnings:
        print("WARNING:", w)


if __name__ == "__main__":
    main()
