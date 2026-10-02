"""Shared constants and openpyxl styling helpers for the report generators in this folder."""
import calendar

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

ADVS = ["TrendiSupply", "TradeExpress", "SmartMediaSolving"]
# Hub invoice type (log "type") -> advertiser. FIND = DE/LT/CZ/SK/LV, ALCANCE = PT/ES (pre-v3 accounts).
TYPE2ADV = {"PL": "TrendiSupply", "FIND": "TrendiSupply", "DE": "TrendiSupply", "LT": "TrendiSupply",
            "CZ": "TrendiSupply", "SK": "TrendiSupply", "LV": "TrendiSupply",
            "HU": "TradeExpress", "ALCANCE": "SmartMediaSolving", "PT": "SmartMediaSolving", "ES": "SmartMediaSolving"}
COUNTRIES = ["PL", "DE", "LT", "CZ", "SK", "LV", "HU", "PT", "ES"]
COUNTRY_NAME = {"PL": "Poland", "DE": "Germany", "LT": "Lithuania", "CZ": "Czech Republic", "SK": "Slovakia",
                "LV": "Latvia", "HU": "Hungary", "PT": "Portugal", "ES": "Spain", "N/A": "No breakdown"}
TITLE_CC = [("Poland", "PL"), ("Hungary", "HU"), ("Germany", "DE"), ("Lithuania", "LT"), ("Czech", "CZ"),
            ("Slovakia", "SK"), ("Latvia", "LV"), ("Portugal", "PT"), ("Spain", "ES")]
NA = "N/A"


def offer_id(raw):
    """AdRice offer id as int; unresolved items (e.g. 'NO_ID') stay text so they remain visible."""
    s = str(raw).strip()
    return int(s) if s.isdigit() else s


def month_label(key):
    y, m = key.split("-")
    return "%s %s" % (calendar.month_name[int(m)], y)


F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=14, bold=True)
f_sec = Font(name=F, size=11, bold=True, color="1F3864")
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="0563C1", underline="single")
f_xs = Font(name=F, size=10, color="008000")
f_sub = Font(name=F, size=10, italic=True, color="595959")
fill_hdr = PatternFill("solid", fgColor="1F3864")
fill_tot = PatternFill("solid", fgColor="D9E1F2")
fill_key = PatternFill("solid", fgColor="E2EFDA")
fill_warn = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
b_all = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00;[Red]-#,##0.00;"-"'   # negatives (credits) with a minus sign, in red
INT = '#,##0;[Red]-#,##0;"-"'
DATE = "dd.mm.yyyy"
PCT = "0.0%"


def put(ws, r, c, v, font=None, fmt=None, fill=None, bold=False, wrap=False, align=None):
    cell = ws.cell(r, c, v)
    cell.font = font if font is not None else (f_bold if bold else f_norm)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return cell


def header(ws, r, c0, labels, height=32):
    for j, h in enumerate(labels):
        c = ws.cell(r, c0 + j, h)
        c.font, c.fill, c.border = f_hdr, fill_hdr, b_all
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = height


def link(ws, r, c, text, sheet, cell):
    x = ws.cell(r, c, text)
    x.font = f_link
    x.hyperlink = Hyperlink(ref=x.coordinate, location="'%s'!%s" % (sheet, cell), display=text)


