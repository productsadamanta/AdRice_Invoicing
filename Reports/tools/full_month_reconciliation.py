#!/usr/bin/env python
"""Full-Month Lead Reconciliation: revenue of ONE lead month per advertiser, split into what was invoiced in
that month's invoices, what was billed later as adjustments, and the prior-month corrections that rode on
the month's invoices. Built from the Invoicing Hub backup only.

Reading the result (per advertiser):
  (A) lead revenue on the month's lead invoices
  (B) adjustments for that month's leads that were billed on LATER invoices (leads approved after close)
  FINAL = A + B   <- revenue of the lead month
  memo  = adjustments for EARLIER months that were billed on this month's invoices
  A + B + memo = the 'Wartosc calkowita' card of the Hub History Center (bridge table on the Summary sheet)

Sheets: Summary, By country, Invoices, Adjustments, Invoice lines, Adjustment lines, By offer, Notes & checks.
All totals are formulas; Notes & checks holds the control sums (all differences must be 0.00).

Usage (from the Invoicing folder):
    python Reports/tools/full_month_reconciliation.py --month 2026-05
Optional: --backup <adrice_backup_*.json> (default: newest in ArchiveBackups), --out <file.xlsx>
"""
import argparse
import collections
import datetime as dt
import glob
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import xlsx_formula_cache  # noqa: E402
from report_common import *  # noqa: E402,F401,F403

TRANSFER_LABELS = {"PRZENIESIENIE_DO_SALDA", "REKONCYLIACJA_DO_SALDA"}
THIS, PRIOR = "This month", "Prior period"   # values of the 'Scope' column on adjustment rows
INV_RE = re.compile(r"INV-?(\d+)$")


def pdt(ts):  # "8.05.2026, 10:55:49"
    d, mo, y, h, mi, s = map(int, re.match(r"(\d+)\.(\d+)\.(\d+),\s*(\d+):(\d+):(\d+)", ts).groups())
    return dt.datetime(y, mo, d, h, mi, s)


def newest_backup(folder):
    files = [f for f in sorted(glob.glob(os.path.join(folder, "adrice_backup_*.json")))
             if not re.search(r"(reconstructed|REPAIRED|merged)", f, re.I)]
    return files[-1]


def log_key(log):
    """Unique key of one Hub log (invoice labels are not unique: 'Google', re-used INV numbers)."""
    return "%s | %s | %s" % (log.get("invoiceNumber") or "-", log["type"], pdt(log["timestamp"]).strftime("%d.%m.%Y %H:%M:%S"))


def is_transfer(log):
    return log.get("invoiceNumber") in TRANSFER_LABELS or "KOREKTA" in (log.get("mode") or "") and not INV_RE.match(log.get("invoiceNumber") or "")


# ---------------------------------------------------------------- extraction
def collect(root, M):
    db = root.get("invoicingDB", root)
    if M not in db:
        raise SystemExit("Month %s is not in the backup." % M)
    month = db[M]
    off_cc = collections.defaultdict(set)
    for cc in COUNTRIES:
        for oid in month.get(cc.lower(), {}):
            off_cc[oid].add(cc)

    def country(it, fallback=None):
        hits = [cc for word, cc in TITLE_CC if re.search(r"\b" + word, it.get("title", ""))]
        if hits:
            return hits[-1]
        c = off_cc.get(str(it["id"]), set())
        if len(c) == 1:
            return next(iter(c))
        return fallback.upper() if fallback and fallback.upper() in COUNTRIES else NA

    def title(it):
        return re.sub(r"^Lead Generation\s+", "", it.get("title", ""))

    warnings = []
    inv_logs = sorted((l for l in month["logs"] if not is_transfer(l)), key=lambda l: (pdt(l["timestamp"]), l["type"]))
    transfers = [l for l in month["logs"] if is_transfer(l)]
    unknown = {l["type"] for l in inv_logs if l["type"] not in TYPE2ADV}
    if unknown:
        warnings.append("Invoice types without an advertiser mapping (left out of the summary): %s" % ", ".join(sorted(unknown)))

    lines = []
    for l in inv_logs:
        for it in l["items"]:
            if it["delta"] != 0:
                lines.append(dict(inv=l["invoiceNumber"], key=log_key(l), date=pdt(l["timestamp"]), acct=l["type"], adv=TYPE2ADV.get(l["type"], "?"),
                                  cc=country(it), oid=offer_id(it["id"]), title=title(it), leads=it["delta"], price=it["price"]))

    tids, seen_t = {}, collections.Counter()
    for l in transfers:  # several logs can share one second (reconciliation batches) -> number them
        base = "TRANSFER %s %s" % (l["type"], pdt(l["timestamp"]).strftime("%d.%m.%Y %H:%M:%S"))
        seen_t[base] += 1
        tids[id(l)] = base if seen_t[base] == 1 else "%s #%d" % (base, seen_t[base])

    def tid(l):
        return tids[id(l)]

    tr_cc, used_tr = {}, set()
    for l in transfers:
        bc = collections.defaultdict(float)
        for it in l["items"]:
            bc[country(it)] += it["delta"] * it["price"]
        tr_cc[tid(l)] = dict(bc)

    def adj_item_lines(a):
        """Lines of one adjustment. Adjustments made by the reconciliation tool carry only an amount (items with
        0 leads), so any amount not explained by leads x price is added as an 'amount-only' line."""
        out, tot = [], 0.0
        for it in a["items"]:
            if it["delta"] == 0:
                continue
            out.append(dict(cc=country(it, a.get("account")), oid=offer_id(it["id"]), title=title(it), leads=it["delta"], price=it["price"], value=None))
            tot += it["delta"] * it["price"]
        resid = a["amount"] - tot
        if abs(resid) > 0.005:
            base = a["items"][0] if a["items"] else {"id": NA, "title": ""}
            out.append(dict(cc=country(base, a.get("account")), oid=offer_id(base["id"]) if a["items"] else NA, leads=None, price=None, value=resid,
                            title="Amount-only correction (no lead quantities in the Hub) - %s" % (title(base) or "n/a")))
        return out

    def per_country(lns):
        per = collections.defaultdict(float)
        for x in lns:
            per[x["cc"]] += x["value"] if x["leads"] is None else x["leads"] * x["price"]
        return per

    def origin(a, lns):
        """Hub log the adjustment came from: same countries and same per-country amounts."""
        per = per_country(lns)
        hits = [t for t, bc in tr_cc.items() if all(c in bc and abs(per[c] - bc[c]) < 1e-9 for c in per)]
        exact = [t for t in hits if abs(sum(tr_cc[t].values()) - a["amount"]) < 1e-9]  # log == adjustment
        if exact:
            pick = ([t for t in exact if t not in used_tr] or exact)[0]
            used_tr.add(pick)  # identical logs (reconciliation batches) are claimed one by one
            return pick
        if hits:  # one log split into several adjustments (e.g. one per country)
            return hits[0]
        if a.get("source") == "reconciliation":
            return "Reconciliation (%s)" % a.get("reconciliationStagingId", "no staging id")
        return "No matching transfer log"

    adj_rows, adj_lines = [], []

    def add_adjustment(a, scope, inv_label, key, date, acct, mode, origin_text=None):
        lns = adj_item_lines(a)
        row = dict(id=a["id"], adv=TYPE2ADV.get(acct, "?"), acct=acct, period=a["month"], scope=scope,
                   origin=origin_text or (origin(a, lns) if scope == THIS else "Hub pending balance created at month close"),
                   amount=a["amount"], inv=inv_label, key=key, date=date, mode=mode,
                   ccs=sorted({x["cc"] for x in lns}, key=lambda c: COUNTRIES.index(c) if c in COUNTRIES else 99))
        adj_rows.append(row)
        for x in lns:
            adj_lines.append(dict(id=a["id"], period=a["month"], scope=scope, adv=row["adv"], inv=inv_label, key=key, date=date, **x))

    for mk in sorted(k for k in db if re.match(r"\d{4}-\d{2}$", k) and k >= M):
        for l in db[mk].get("logs", []):
            for a in l.get("adjustments", []):
                if a["month"] == M:
                    scope = THIS
                elif mk == M and a["month"] < M:
                    scope = PRIOR
                else:
                    continue
                add_adjustment(a, scope, l["invoiceNumber"], log_key(l), pdt(l["timestamp"]), l["type"], l["mode"])
    for a in root.get("pendingAdjustments", []):  # not yet billed on any invoice
        if a.get("month") == M and a.get("status") != "used":
            add_adjustment(a, THIS, "(pending)", "(pending)", None, a["account"].upper(), "", origin_text="Pending - not yet billed")
            warnings.append("Adjustment %s (EUR %.2f) for %s is still pending - it is counted in (B) but not billed yet." % (a["id"], a["amount"], month_label(M)))
    adj_rows.sort(key=lambda r: (r["scope"] != PRIOR, r["date"] or dt.datetime.max, r["acct"], r["id"]))

    # ---- facts for the notes
    all_nums = set()
    for mk, m in db.items():
        if isinstance(m, dict) and "logs" in m:
            for l in m["logs"]:
                mm = INV_RE.match(l.get("invoiceNumber") or "")
                if mm:
                    all_nums.add(int(mm.group(1)))
    nums = sorted(int(INV_RE.match(l["invoiceNumber"]).group(1)) for l in inv_logs if INV_RE.match(l["invoiceNumber"] or ""))
    gaps = [n for n in range(nums[0], nums[-1] + 1) if n not in all_nums] if nums else []
    # lead counter chain: each invoice must start where the previous one for the same offer ended
    chain, breaks, nlines, last = {}, 0, 0, {}
    for l in inv_logs:
        for it in l["items"]:
            if "prev" not in it:
                continue
            k = (country(it), str(it["id"]))
            nlines += 1
            if it["curr"] - it["prev"] != it["delta"] or (k in last and last[k] != it["prev"]):
                breaks += 1
            last[k] = it["curr"]
    trl = collections.defaultdict(int)
    for l in transfers:
        for it in l["items"]:
            trl[(country(it), str(it["id"]))] += it["delta"]
    same = plus_transfer = other = 0
    for (cc, oid), cur in last.items():
        st = month.get(cc.lower(), {}).get(oid, 0)
        if st == cur:
            same += 1
        elif st == cur + trl.get((cc, oid), 0):
            plus_transfer += 1
        else:
            other += 1
    facts = dict(n_inv=len(inv_logs), d0=min(pdt(l["timestamp"]) for l in inv_logs), d1=max(pdt(l["timestamp"]) for l in inv_logs),
                 num0=nums[0] if nums else None, num1=nums[-1] if nums else None, gaps=gaps, transfers=[(tid(l), l["type"]) for l in transfers],
                 nlines=nlines, breaks=breaks, same=same, plus_transfer=plus_transfer, other=other, tr_cc=tr_cc, tid=tid)
    return dict(lines=lines, adj_rows=adj_rows, adj_lines=adj_lines, inv_logs=inv_logs, transfers=transfers, facts=facts, warnings=warnings)


# ---------------------------------------------------------------- workbook
def build(root, M, backup_name):
    data = collect(root, M)
    lines, adj_rows, adj_lines, inv_logs, F_ = data["lines"], data["adj_rows"], data["adj_lines"], data["inv_logs"], data["facts"]
    ML = month_label(M)
    wb = Workbook()

    # ---- Invoice lines
    wsL = wb.active
    wsL.title = "Invoice lines"
    header(wsL, 1, 1, ["Invoice no. (Hub)", "Invoice date", "Hub account", "Advertiser", "Country", "Offer ID", "Offer", "Leads", "Price (€)", "Amount (€)", "Invoice key (unique)"], 42)
    for k, w in enumerate([16, 17, 12, 18, 10, 9, 70, 9, 10, 13, 34], 1):
        wsL.column_dimensions[L(k)].width = w
    for i, x in enumerate(lines, start=2):
        put(wsL, i, 1, x["inv"]); put(wsL, i, 2, x["date"], fmt="dd.mm.yyyy hh:mm"); put(wsL, i, 3, x["acct"]); put(wsL, i, 4, x["adv"])
        put(wsL, i, 5, x["cc"]); put(wsL, i, 6, x["oid"]); put(wsL, i, 7, x["title"])
        put(wsL, i, 8, x["leads"], font=f_in, fmt=INT); put(wsL, i, 9, x["price"], font=f_in, fmt=EUR); put(wsL, i, 10, "=H%d*I%d" % (i, i), fmt=EUR); put(wsL, i, 11, x["key"])
    NL = len(lines) + 1
    wsL.freeze_panes = "A2"
    wsL.auto_filter.ref = "A1:K%d" % NL
    LR = lambda c: "'Invoice lines'!$%s$2:$%s$%d" % (c, c, NL)  # noqa: E731

    # ---- Adjustment lines
    wsAL = wb.create_sheet("Adjustment lines")
    header(wsAL, 1, 1, ["Adjustment ID", "Relates to (month)", "Advertiser", "Country", "Offer ID", "Offer", "Leads", "Price (€)", "Amount (€)",
                        "Billed on invoice (Hub label)", "Invoice date", "Scope", "Invoice key (unique)"], 42)
    for k, w in enumerate([34, 16, 18, 10, 9, 70, 9, 10, 13, 18, 17, 13, 34], 1):
        wsAL.column_dimensions[L(k)].width = w
    for i, x in enumerate(adj_lines, start=2):
        put(wsAL, i, 1, x["id"]); put(wsAL, i, 2, month_label(x["period"])); put(wsAL, i, 3, x["adv"]); put(wsAL, i, 4, x["cc"])
        put(wsAL, i, 5, x["oid"]); put(wsAL, i, 6, x["title"])
        if x["leads"] is None:  # amount-only correction from the reconciliation tool: value comes straight from the Hub adjustment
            put(wsAL, i, 9, x["value"], font=f_in, fmt=EUR, fill=fill_warn)
            wsAL.cell(i, 6).fill = fill_warn
        else:
            put(wsAL, i, 7, x["leads"], font=f_in, fmt=INT); put(wsAL, i, 8, x["price"], font=f_in, fmt=EUR); put(wsAL, i, 9, "=G%d*H%d" % (i, i), fmt=EUR)
        put(wsAL, i, 10, x["inv"]); put(wsAL, i, 11, x["date"], fmt="dd.mm.yyyy hh:mm"); put(wsAL, i, 12, x["scope"]); put(wsAL, i, 13, x["key"])
    NAL = max(len(adj_lines) + 1, 2)
    wsAL.freeze_panes = "A2"
    wsAL.auto_filter.ref = "A1:M%d" % NAL
    AL = lambda c: "'Adjustment lines'!$%s$2:$%s$%d" % (c, c, NAL)  # noqa: E731

    # ---- Invoices
    wsI = wb.create_sheet("Invoices")
    header(wsI, 1, 1, ["Advertiser", "Hub account", "Invoice no. (Hub)", "Invoice date", "Mode", "Countries on invoice", "Leads", "Lead lines (€)",
                       "Prior-period adjustments on invoice (€)", "Invoice total as issued (€)", "Invoice key (unique)"], 42)
    for k, w in enumerate([18, 12, 16, 17, 20, 26, 9, 15, 17, 17, 34], 1):
        wsI.column_dimensions[L(k)].width = w
    for i, l in enumerate(inv_logs, start=2):
        ccs = sorted({x["cc"] for x in lines if x["key"] == log_key(l)}, key=lambda c: COUNTRIES.index(c) if c in COUNTRIES else 99)
        put(wsI, i, 1, TYPE2ADV.get(l["type"], "?")); put(wsI, i, 2, l["type"]); put(wsI, i, 3, l["invoiceNumber"]); put(wsI, i, 4, pdt(l["timestamp"]), fmt="dd.mm.yyyy hh:mm")
        put(wsI, i, 5, l["mode"]); put(wsI, i, 6, ", ".join(ccs))
        put(wsI, i, 7, "=SUMIFS(%s,%s,$K%d)" % (LR("H"), LR("K"), i), fmt=INT)
        put(wsI, i, 8, "=SUMIFS(%s,%s,$K%d)" % (LR("J"), LR("K"), i), fmt=EUR)
        put(wsI, i, 9, '=SUMIFS(%s,%s,$K%d,%s,"%s")' % (AL("I"), AL("M"), i, AL("L"), PRIOR), fmt=EUR)
        put(wsI, i, 10, "=H%d+I%d" % (i, i), fmt=EUR); put(wsI, i, 11, log_key(l))
    NI = len(inv_logs) + 1
    tI = NI + 1
    put(wsI, tI, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in list(range(2, 7)) + [11]:
        put(wsI, tI, c, None, fill=fill_tot)
    for c, fm in ((7, INT), (8, EUR), (9, EUR), (10, EUR)):
        put(wsI, tI, c, "=SUM(%s2:%s%d)" % (L(c), L(c), NI), fmt=fm, bold=True, fill=fill_tot)
    wsI.freeze_panes = "A2"
    wsI.auto_filter.ref = "A1:K%d" % NI
    IR = lambda c: "Invoices!$%s$2:$%s$%d" % (c, c, NI)  # noqa: E731

    # ---- Adjustments
    wsA = wb.create_sheet("Adjustments")
    header(wsA, 1, 1, ["Adjustment ID", "Advertiser", "Hub account", "Countries", "Relates to (month)", "Scope", "Origin of the adjustment (Hub log)", "Amount per Hub (€)",
                       "Σ of lines (€)", "Difference (€)", "Billed on invoice (Hub label)", "Invoice date", "Invoice mode", "Real invoice no. (optional)"], 42)
    for k, w in enumerate([34, 18, 11, 16, 14, 13, 50, 15, 14, 12, 18, 17, 22, 24], 1):
        wsA.column_dimensions[L(k)].width = w
    for i, r in enumerate(adj_rows, start=2):
        put(wsA, i, 1, r["id"]); put(wsA, i, 2, r["adv"]); put(wsA, i, 3, r["acct"]); put(wsA, i, 4, ", ".join(r["ccs"])); put(wsA, i, 5, month_label(r["period"]))
        put(wsA, i, 6, r["scope"]); put(wsA, i, 7, r["origin"]); put(wsA, i, 8, r["amount"], font=f_in, fmt=EUR)
        put(wsA, i, 9, "=SUMIFS(%s,%s,A%d)" % (AL("I"), AL("A"), i), fmt=EUR); put(wsA, i, 10, "=H%d-I%d" % (i, i), fmt=EUR)
        put(wsA, i, 11, r["inv"]); put(wsA, i, 12, r["date"], fmt="dd.mm.yyyy hh:mm"); put(wsA, i, 13, r["mode"])
        put(wsA, i, 14, None, fill=fill_warn if not INV_RE.match(r["inv"]) else None)
    NA_ = max(len(adj_rows) + 1, 2)
    tA = NA_ + 1
    put(wsA, tA, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in range(2, 8):
        put(wsA, tA, c, None, fill=fill_tot)
    for c in (8, 9, 10):
        put(wsA, tA, c, "=SUM(%s2:%s%d)" % (L(c), L(c), NA_), fmt=EUR, bold=True, fill=fill_tot)
    for c in range(11, 15):
        put(wsA, tA, c, None, fill=fill_tot)
    wsA.freeze_panes = "A2"
    wsA.auto_filter.ref = "A1:N%d" % NA_
    AR = lambda c: "Adjustments!$%s$2:$%s$%d" % (c, c, NA_)  # noqa: E731

    # ---- By country
    wsC = wb.create_sheet("By country")
    put(wsC, 1, 1, "%s - revenue by advertiser and country" % ML, font=f_title)
    put(wsC, 2, 1, "EUR, from the Invoicing Hub logs. FINAL = lead revenue on %s invoices + %s adjustments billed later. Prior-period adjustments are a memo only." % (ML, ML), font=f_sub)
    header(wsC, 4, 1, ["Advertiser", "Country", "Leads on %s invoices" % ML, "(A) Lead revenue on %s invoices (€)" % ML, "Leads in %s adjustments" % ML,
                       "(B) %s adjustments billed later (€)" % ML, "FINAL reconciled %s revenue (A+B) (€)" % ML, "Memo: prior-period adjustments on %s invoices (€)" % ML], 56)
    for k, w in enumerate([20, 20, 13, 19, 13, 19, 20, 19], 1):
        wsC.column_dimensions[L(k)].width = w
    present = {(x["adv"], x["cc"]) for x in lines} | {(x["adv"], x["cc"]) for x in adj_lines}
    r = 5
    subs, firsts = {}, {}
    for adv in ADVS:
        firsts[adv] = r
        for cc in [c for c in COUNTRIES + [NA] if (adv, c) in present]:
            key = '"%s"' % cc
            put(wsC, r, 1, adv); put(wsC, r, 2, "%s - %s" % (cc, COUNTRY_NAME.get(cc, cc)))
            put(wsC, r, 3, "=SUMIFS(%s,%s,$A%d,%s,%s)" % (LR("H"), LR("D"), r, LR("E"), key), fmt=INT)
            put(wsC, r, 4, "=SUMIFS(%s,%s,$A%d,%s,%s)" % (LR("J"), LR("D"), r, LR("E"), key), fmt=EUR)
            put(wsC, r, 5, '=SUMIFS(%s,%s,$A%d,%s,%s,%s,"%s")' % (AL("G"), AL("C"), r, AL("D"), key, AL("L"), THIS), fmt=INT)
            put(wsC, r, 6, '=SUMIFS(%s,%s,$A%d,%s,%s,%s,"%s")' % (AL("I"), AL("C"), r, AL("D"), key, AL("L"), THIS), fmt=EUR)
            put(wsC, r, 7, "=D%d+F%d" % (r, r), fmt=EUR, bold=True)
            put(wsC, r, 8, '=SUMIFS(%s,%s,$A%d,%s,%s,%s,"%s")' % (AL("I"), AL("C"), r, AL("D"), key, AL("L"), PRIOR), fmt=EUR)
            r += 1
        put(wsC, r, 1, adv + " - subtotal", bold=True, fill=fill_tot)
        put(wsC, r, 2, None, fill=fill_tot)
        for c, fm in ((3, INT), (4, EUR), (5, INT), (6, EUR), (7, EUR), (8, EUR)):
            put(wsC, r, c, "=SUM(%s%d:%s%d)" % (L(c), firsts[adv], L(c), r - 1), fmt=fm, bold=True, fill=fill_tot)
        subs[adv] = r
        r += 1
    put(wsC, r, 1, "GRAND TOTAL", bold=True, fill=fill_key)
    put(wsC, r, 2, None, fill=fill_key)
    for c, fm in ((3, INT), (4, EUR), (5, INT), (6, EUR), (7, EUR), (8, EUR)):
        put(wsC, r, c, "=" + "+".join("%s%d" % (L(c), subs[a]) for a in ADVS), fmt=fm, bold=True, fill=fill_key)
    C_GRAND = r
    wsC.freeze_panes = "A5"

    # ---- By offer
    wsO = wb.create_sheet("By offer")
    put(wsO, 1, 1, "%s - revenue by offer (supporting detail)" % ML, font=f_title)
    put(wsO, 2, 1, "One row per advertiser / country / offer. Prices can change during the month, so the average price is shown instead of a single rate.", font=f_sub)
    header(wsO, 4, 1, ["Advertiser", "Country", "Offer ID", "Offer", "Leads on %s invoices" % ML, "(A) Lead revenue on %s invoices (€)" % ML, "Leads in %s adjustments" % ML,
                       "(B) %s adjustments billed later (€)" % ML, "Total leads", "FINAL %s revenue (A+B) (€)" % ML, "Avg price (€/lead)"], 56)
    for k, w in enumerate([19, 10, 9, 70, 12, 18, 12, 18, 10, 17, 12], 1):
        wsO.column_dimensions[L(k)].width = w
    offers, amt = {}, collections.defaultdict(float)
    for x in lines:
        offers.setdefault((x["adv"], x["cc"], x["oid"]), x["title"]); amt[(x["adv"], x["cc"], x["oid"])] += x["leads"] * x["price"]
    for x in adj_lines:
        if x["scope"] == THIS:
            offers.setdefault((x["adv"], x["cc"], x["oid"]), x["title"]); amt[(x["adv"], x["cc"], x["oid"])] += x["value"] if x["leads"] is None else x["leads"] * x["price"]
    okeys = sorted(offers, key=lambda k: (ADVS.index(k[0]) if k[0] in ADVS else 99, COUNTRIES.index(k[1]) if k[1] in COUNTRIES else 99, -amt[k], str(k[2])))
    for i, k in enumerate(okeys, start=5):
        adv, cc, oid = k
        put(wsO, i, 1, adv); put(wsO, i, 2, cc); put(wsO, i, 3, oid); put(wsO, i, 4, offers[k])
        put(wsO, i, 5, "=SUMIFS(%s,%s,$A%d,%s,$B%d,%s,$C%d)" % (LR("H"), LR("D"), i, LR("E"), i, LR("F"), i), fmt=INT)
        put(wsO, i, 6, "=SUMIFS(%s,%s,$A%d,%s,$B%d,%s,$C%d)" % (LR("J"), LR("D"), i, LR("E"), i, LR("F"), i), fmt=EUR)
        put(wsO, i, 7, '=SUMIFS(%s,%s,$A%d,%s,$B%d,%s,$C%d,%s,"%s")' % (AL("G"), AL("C"), i, AL("D"), i, AL("E"), i, AL("L"), THIS), fmt=INT)
        put(wsO, i, 8, '=SUMIFS(%s,%s,$A%d,%s,$B%d,%s,$C%d,%s,"%s")' % (AL("I"), AL("C"), i, AL("D"), i, AL("E"), i, AL("L"), THIS), fmt=EUR)
        put(wsO, i, 9, "=E%d+G%d" % (i, i), fmt=INT); put(wsO, i, 10, "=F%d+H%d" % (i, i), fmt=EUR, bold=True)
        put(wsO, i, 11, "=IF(I%d=0,0,J%d/I%d)" % (i, i, i), fmt=EUR)
    NO = 4 + len(okeys)
    tO = NO + 1
    put(wsO, tO, 1, "TOTAL", bold=True, fill=fill_tot)
    for c in (2, 3, 4, 11):
        put(wsO, tO, c, None, fill=fill_tot)
    for c, fm in ((5, INT), (6, EUR), (7, INT), (8, EUR), (9, INT), (10, EUR)):
        put(wsO, tO, c, "=SUM(%s5:%s%d)" % (L(c), L(c), NO), fmt=fm, bold=True, fill=fill_tot)
    wsO.freeze_panes = "E5"
    wsO.auto_filter.ref = "A4:K%d" % NO

    # ---- Summary
    wsS = wb.create_sheet("Summary", 0)
    put(wsS, 1, 1, "Revenue reconciliation - %s" % ML, font=f_title)
    put(wsS, 2, 1, "Source: AdRice Invoicing Hub (invoice history + balance adjustments), backup %s. Amounts in EUR, net as recorded in the Hub." % backup_name, font=f_sub)
    header(wsS, 4, 1, ["Advertiser", "Hub accounts / countries", "# lead invoices in %s" % ML, "Leads on %s invoices" % ML, "(A) Lead revenue on %s invoices (€)" % ML,
                       "# adjustments for %s billed later" % ML, "Leads in %s adjustments" % ML, "(B) %s adjustments billed later (€)" % ML,
                       "FINAL reconciled %s revenue (A+B) (€)" % ML], 56)
    for k, w in enumerate([20, 30, 16, 16, 19, 16, 16, 19, 21], 1):
        wsS.column_dimensions[L(k)].width = w
    cov = {}
    for adv in ADVS:
        ccs = [c for c in COUNTRIES if (adv, c) in present]
        accts = sorted({x["acct"] for x in lines if x["adv"] == adv})
        cov[adv] = "%s  (Hub: %s)" % (", ".join(ccs) if ccs else "-", ", ".join(accts) if accts else "-")
    for k, adv in enumerate(ADVS):
        r = 5 + k
        put(wsS, r, 1, adv, bold=True); put(wsS, r, 2, cov[adv], wrap=True); wsS.row_dimensions[r].height = 28
        put(wsS, r, 3, "=COUNTIFS(%s,$A%d)" % (IR("A"), r), fmt=INT)
        put(wsS, r, 4, "=SUMIFS(%s,%s,$A%d)" % (LR("H"), LR("D"), r), fmt=INT)
        put(wsS, r, 5, "=SUMIFS(%s,%s,$A%d)" % (LR("J"), LR("D"), r), fmt=EUR)
        put(wsS, r, 6, '=COUNTIFS(%s,$A%d,%s,"%s")' % (AR("B"), r, AR("F"), THIS), fmt=INT)
        put(wsS, r, 7, '=SUMIFS(%s,%s,$A%d,%s,"%s")' % (AL("G"), AL("C"), r, AL("L"), THIS), fmt=INT)
        put(wsS, r, 8, '=SUMIFS(%s,%s,$A%d,%s,"%s")' % (AL("I"), AL("C"), r, AL("L"), THIS), fmt=EUR)
        put(wsS, r, 9, "=E%d+H%d" % (r, r), fmt=EUR, bold=True, fill=fill_key)
    put(wsS, 8, 1, "TOTAL", bold=True, fill=fill_tot); put(wsS, 8, 2, None, fill=fill_tot)
    for c, fm in ((3, INT), (4, INT), (5, EUR), (6, INT), (7, INT), (8, EUR), (9, EUR)):
        put(wsS, 8, c, "=SUM(%s5:%s7)" % (L(c), L(c)), fmt=fm, bold=True, fill=fill_tot)

    put(wsS, 10, 1, "Bridge to the Invoicing Hub History Center (company tab, month %s)" % ML, font=f_bold)
    header(wsS, 11, 1, ["Advertiser", "FINAL reconciled %s revenue (A+B) (€)" % ML, "Prior-period adjustments billed on %s invoices (€)" % ML,
                        "History Center 'Wartość całkowita' (€) = FINAL + prior-period adj.", "Total %s leads = History Center 'Lidów w tym miesiącu'" % ML,
                        "Memo: %s invoices as issued (A + prior-period adj.) (€)" % ML, "Check: Σ Invoices sheet (€)", "Difference (€)"], 70)
    for k, adv in enumerate(ADVS):
        r, m = 12 + k, 5 + k
        put(wsS, r, 1, adv, bold=True); put(wsS, r, 2, "=I%d" % m, fmt=EUR)
        put(wsS, r, 3, '=SUMIFS(%s,%s,$A%d,%s,"%s")' % (AL("I"), AL("C"), r, AL("L"), PRIOR), fmt=EUR)
        put(wsS, r, 4, "=B%d+C%d" % (r, r), fmt=EUR, bold=True); put(wsS, r, 5, "=D%d+G%d" % (m, m), fmt=INT, bold=True)
        put(wsS, r, 6, "=E%d+C%d" % (m, r), fmt=EUR); put(wsS, r, 7, "=SUMIFS(%s,%s,$A%d)" % (IR("J"), IR("A"), r), fmt=EUR); put(wsS, r, 8, "=F%d-G%d" % (r, r), fmt=EUR)
    put(wsS, 15, 1, "TOTAL", bold=True, fill=fill_tot)
    for c, fm in ((2, EUR), (3, EUR), (4, EUR), (5, INT), (6, EUR), (7, EUR), (8, EUR)):
        put(wsS, 15, c, "=SUM(%s12:%s14)" % (L(c), L(c)), fmt=fm, bold=True, fill=fill_tot)
    put(wsS, 17, 1, "How to read this report", font=f_bold)
    how = [
        "Revenue entry = %s revenue per advertiser. Column I (green) is the lead-month figure: all leads of %s, whenever they were invoiced." % (ML, ML),
        "(A) = leads x price on every lead invoice the Hub recorded for %s (%s - %s). Detail: sheets 'Invoices' (one row per invoice) and 'Invoice lines' (one row per offer)." % (ML, F_["d0"].strftime("%d.%m"), F_["d1"].strftime("%d.%m.%Y")),
        "(B) = %s leads approved after the %s invoices; they were moved to the Hub balance and billed as adjustment lines on later invoices. Detail: sheets 'Adjustments' (links each to its source log and to the invoice that carried it) and 'Adjustment lines'." % (ML, ML),
        "The History Center card shows (A) + (B) + the corrections for earlier months that were billed on the %s invoices. Those are earlier-month leads, so the lead-month figure (column I) excludes them; the bridge table above shows both numbers side by side." % ML,
        "Per-country split: sheet 'By country'. Per-offer split: sheet 'By offer'. Control totals and notes: sheet 'Notes & checks'.",
        "Blue font = values taken from the Hub; black = formulas. Yellow cells (sheet 'Adjustments', last column) = optional real invoice numbers.",
    ]
    for i, t in enumerate(how, start=18):
        put(wsS, i, 1, "•"); put(wsS, i, 2, t, wrap=True)
        wsS.merge_cells(start_row=i, start_column=2, end_row=i, end_column=9)
        wsS.row_dimensions[i].height = 44 if len(t) > 190 else 30
    wsS.freeze_panes = "A5"

    # ---- Notes & checks
    wsN = wb.create_sheet("Notes & checks")
    wsN.column_dimensions["A"].width = 64
    for c in "BCDE":
        wsN.column_dimensions[c].width = 19
    wsN.column_dimensions["F"].width = 60
    put(wsN, 1, 1, "Control checks (all differences should be 0.00)", font=f_title)
    header(wsN, 3, 1, ["Check", "Value 1 (€)", "Value 2 (€)", "Difference (€)"])
    checks = [
        ("Σ Invoice lines = Σ Invoices sheet 'Lead lines'", "=SUM(%s)" % LR("J"), "=Invoices!H%d" % tI),
        ("Σ Summary (A) = Σ Invoice lines", "=Summary!E8", "=SUM(%s)" % LR("J")),
        ("Σ Summary FINAL = By country grand total", "=Summary!I8", "='By country'!G%d" % C_GRAND),
        ("Σ Summary FINAL = By offer total", "=Summary!I8", "='By offer'!J%d" % tO),
        ("Σ Summary total leads = By country leads (invoices + adjustments)", "=Summary!D8+Summary!G8", "='By country'!C%d+'By country'!E%d" % (C_GRAND, C_GRAND)),
        ("Σ 'invoices as issued' (bridge) = Σ Invoices sheet total", "=Summary!F15", "=Invoices!J%d" % tI),
        ("Σ adjustment amounts per Hub = Σ adjustment lines", "=Adjustments!H%d" % tA, "=SUM(%s)" % AL("I")),
    ]
    for k, (t, f1, f2) in enumerate(checks, start=4):
        put(wsN, k, 1, t); put(wsN, k, 2, f1, fmt=EUR); put(wsN, k, 3, f2, fmt=EUR); put(wsN, k, 4, "=B%d-C%d" % (k, k), fmt=EUR)
    r0 = 4 + len(checks) + 1
    put(wsN, r0, 1, "Origin tie-out: %s balance-transfer logs in the Hub vs adjustments billed later" % ML, font=f_bold)
    header(wsN, r0 + 1, 1, ["Hub balance-transfer log (leads moved to balance)", "Amount per log (€)", "Σ adjustments billed (€)", "Difference (€)"])
    r = r0 + 2
    ftr = r
    for t, _ in F_["transfers"]:
        put(wsN, r, 1, t); put(wsN, r, 2, sum(F_["tr_cc"][t].values()), font=f_in, fmt=EUR)
        put(wsN, r, 3, '=SUMIFS(%s,%s,A%d,%s,"%s")' % (AR("H"), AR("G"), r, AR("F"), THIS), fmt=EUR); put(wsN, r, 4, "=B%d-C%d" % (r, r), fmt=EUR)
        r += 1
    log_end = r - 1
    put(wsN, r, 1, "TOTAL (logs)", bold=True, fill=fill_tot)
    for c in (2, 3, 4):
        put(wsN, r, c, "=SUM(%s%d:%s%d)" % (L(c), ftr, L(c), log_end), fmt=EUR, bold=True, fill=fill_tot)
    r += 1
    put(wsN, r, 1, "Adjustments from another origin (reconciliation tool, pending, no matching log) - information only")
    put(wsN, r, 3, '=SUMIFS(%s,%s,"%s")-C%d' % (AR("H"), AR("F"), THIS, r - 1), fmt=EUR)
    r += 2
    put(wsN, r, 1, "Scope, method and open points", font=f_title)
    r += 1
    labels_non_inv = sorted({x["inv"] for x in adj_rows if x["scope"] == THIS and not INV_RE.match(x["inv"])})
    prior = collections.OrderedDict()
    for x in adj_rows:
        if x["scope"] == PRIOR:
            prior.setdefault(x["inv"], 0.0)
            prior[x["inv"]] += x["amount"]
    later_inv = sorted({x["inv"] for x in adj_rows if x["scope"] == THIS and INV_RE.match(x["inv"])})
    notes = [
        ("Source", "Invoicing Hub backup %s." % backup_name),
        ("Scope", "All leads the Hub attributes to %s: %d lead invoices dated %s-%s (%s) plus the %s adjustments (%d balance-transfer log(s)%s) that were billed as adjustment lines on later invoices%s." % (
            ML, F_["n_inv"], F_["d0"].strftime("%d.%m"), F_["d1"].strftime("%d.%m.%Y"),
            "INV-%03d...INV-%03d" % (F_["num0"], F_["num1"]) if F_["num0"] else "no numbered invoices", ML, len(F_["transfers"]),
            ": " + ", ".join(sorted({t.split(" ", 2)[2][:10] for t, _ in F_["transfers"]})) if F_["transfers"] else "",
            " (" + ", ".join(later_inv[:6]) + (", ..." if len(later_inv) > 6 else "") + ")" if later_inv else "")),
        ("Amounts", "Invoice amount = Σ(leads x price) + adjustments, computed the same way as the Hub's History Center. Currency EUR. The Hub stores net amounts only; VAT is not included."),
        ("Advertiser mapping", "By Hub account: TrendiSupply = PL and FIND (DE/LT/CZ/SK/LV), TradeExpress = HU, SmartMediaSolving = ALCANCE (PT/ES). This mapping holds from April 2026, when EuroFlex was replaced by TrendiSupply."),
    ]
    if F_["gaps"]:
        notes.append(("Invoice numbers", "Numbers are the labels typed into the Hub when the invoice was generated. %d numbers in the range INV-%03d...%03d have no record in the Hub: %s. They may belong to other companies or to documents issued outside the Hub - check them against the QuickBooks statements (a statement can list an invoice the Hub does not know)." % (
            len(F_["gaps"]), F_["num0"], F_["num1"], ", ".join("%03d" % n for n in F_["gaps"]))))
    if labels_non_inv:
        n_lab = sum(1 for x in adj_rows if x["scope"] == THIS and x["inv"] in labels_non_inv)
        amt_lab = sum(x["amount"] for x in adj_rows if x["scope"] == THIS and x["inv"] in labels_non_inv)
        notes.append(("Invoices without a number", "%d %s adjustments (EUR %s) were billed on invoices labelled %s in the Hub instead of a number. Apart from those adjustment lines these invoices contain leads of the following month, so they matter only as the document that carried the corrections. A real number can be added in 'Adjustments' (yellow column); it does not affect any amount." % (
            n_lab, ML, "{:,.2f}".format(amt_lab), ", ".join("'%s'" % x for x in labels_non_inv))))
    if prior:
        notes.append(("Prior-period adjustments", "%s include adjustments for earlier months (%s; total EUR %s). They are part of those invoices' totals but relate to earlier leads, so they are excluded from the lead-month FINAL figure and shown separately in the bridge table. If the ledger books revenue by invoice date instead, use the bridge column 'invoices as issued'." % (
            ", ".join(prior), ", ".join("%s EUR %s" % (k, "{:,.2f}".format(v).rstrip("0").rstrip(".")) for k, v in prior.items()), "{:,.2f}".format(sum(prior.values())))))
    amount_only = [x for x in adj_lines if x["leads"] is None and x["scope"] == THIS]
    if amount_only:
        notes.append(("Amount-only corrections", "%d correction line(s) of %s (EUR %s in total) come from the Hub reconciliation tool, which records only an amount and no lead quantities. They are included in (B) and in the offer totals with 0 leads (yellow lines in 'Adjustment lines')." % (
            len(amount_only), ML, "{:,.2f}".format(sum(x["value"] for x in amount_only)))))
    notes.append(("Lead counter integrity", "Across the %s invoices the lead counter chain is %s (%d lines checked, %d break(s)). The last counter on the invoices equals the Hub's month counter for %d offer/country combinations; for %d the Hub counter is higher by exactly the leads moved to balance (reported here as adjustments); %d differ for another reason%s." % (
        ML, "continuous (each invoice starts where the previous one ended)" if F_["breaks"] == 0 else "NOT continuous", F_["nlines"], F_["breaks"], F_["same"], F_["plus_transfer"], F_["other"],
        "" if F_["other"] == 0 else " - investigate")))
    notes.append(("History Center", "The company tabs of the Hub History Center add up every log of the month: invoice lines, the balance-transfer logs (leads moved to balance) and the adjustments recorded on invoices. 'Wartość całkowita' = (A) + (B) + prior-period adjustments; 'Lidów w tym miesiącu' = leads on invoices + leads moved to balance. See the bridge table on the Summary sheet."))
    notes.append(("Not covered", "Amounts actually paid or booked in the ledger, VAT, and documents issued outside the Hub. This report ties revenue to Hub invoices and adjustments only."))
    for w in data["warnings"]:
        notes.append(("ATTENTION", w))
    for t, txt in notes:
        put(wsN, r, 1, t, bold=True); wsN.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
        put(wsN, r, 2, txt, wrap=True, fill=fill_warn if t == "ATTENTION" else None)
        wsN.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        wsN.row_dimensions[r].height = max(30, 14 * (1 + len(txt) // 110))
        r += 1

    wb._sheets = [wb[n] for n in ("Summary", "By country", "Invoices", "Adjustments", "Invoice lines", "Adjustment lines", "By offer", "Notes & checks")]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    return wb, data


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--month", required=True, help="Lead month to reconcile, YYYY-MM")
    ap.add_argument("--backup", help="Hub backup JSON (default: newest in ArchiveBackups)")
    ap.add_argument("--out", help="Output .xlsx")
    a = ap.parse_args()
    here = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.dirname(os.path.dirname(here))
    backup = a.backup or newest_backup(os.path.join(root_dir, "ArchiveBackups"))
    out = a.out or os.path.join(root_dir, "Reports", "Full_Month_Reconciliation_%s.xlsx" % a.month)
    root = json.load(open(backup, encoding="utf-8"))
    wb, data = build(root, a.month, os.path.basename(backup))
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    written, total = xlsx_formula_cache.add_cached_values(out)
    print("saved", out)
    print("backup:", backup)
    print("cached values: %d of %d formulas" % (written, total))
    for w in data["warnings"]:
        print("WARNING:", w)


if __name__ == "__main__":
    main()
